from collections import Counter
from datetime import date, datetime
from io import BytesIO

from django.db.models import Q
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import HIRA, HIRAAction, HIRAHazard, RiskMatrixLevel
from .source_adapters import describe_source_object, get_adapter


RISK_LEVEL_ROWS = [
    ("1-4", "Low", "Risk is generally acceptable with controls maintained", "Monitor and maintain controls"),
    ("5-9", "Medium", "Additional improvement may be required", "Improve controls where reasonably practicable"),
    ("10-16", "High", "Significant risk requiring prompt action", "Implement additional controls and management attention"),
    ("17-25", "Critical", "Unacceptable risk", "Stop/restrict activity until adequate controls are implemented"),
]


def _style_sheet(ws):
    thin = Side(style="thin", color="D9E2EC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1


def _autosize(ws, max_width=45):
    for column in ws.columns:
        letter = get_column_letter(column[0].column)
        longest = max((len(str(cell.value)) for cell in column if cell.value is not None), default=10)
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), max_width)


def _risk_fill(value):
    return {
        "Critical": "FFC7CE",
        "High": "FFD9B3",
        "Medium": "FFEB9C",
        "Low": "C6EFCE",
    }.get(value)


def _configured_risk_level_rows():
    rows = list(
        RiskMatrixLevel.objects.filter(
            risk_matrix__is_active=True,
            is_active=True,
        ).select_related("risk_matrix").order_by(
            "-risk_matrix__created_at",
            "display_order",
            "min_score",
        ).values_list(
            "min_score",
            "max_score",
            "risk_level",
            "description",
        )
    )
    if rows:
        return [
            (f"{min_score}-{max_score}", risk_level, description, "")
            for min_score, max_score, risk_level, description in rows
        ]
    return RISK_LEVEL_ROWS


def _status_text(value):
    return dict(HIRA.STATUS_CHOICES).get(value, value)


def _action_status_text(value):
    return dict(HIRAAction.STATUS_CHOICES).get(value, value)


def filter_hira_queryset(request, form=None):
    qs = HIRA.objects.select_related(
        "plant",
        "department",
        "module",
        "prepared_by",
        "reviewed_by",
        "approved_by",
    ).prefetch_related("hazards__actions")
    user = request.user
    is_admin = user.is_superuser or getattr(user, "is_admin_user", False)
    if not is_admin and hasattr(user, "get_all_plants"):
        plants = user.get_all_plants()
        qs = qs.filter(plant__in=plants) if plants else qs.none()

    data = form.cleaned_data if form and form.is_valid() else request.GET
    if data.get("plant"):
        qs = qs.filter(plant=data["plant"] if form else data["plant"])
    if data.get("department"):
        qs = qs.filter(department=data["department"] if form else data["department"])
    if data.get("module"):
        qs = qs.filter(module=data["module"] if form else data["module"])
    if data.get("process"):
        qs = qs.filter(process__icontains=data["process"])
    if data.get("activity"):
        qs = qs.filter(hazards__activity__icontains=data["activity"])
    if data.get("status"):
        qs = qs.filter(status=data["status"])
    if data.get("source_module"):
        adapter = get_adapter(data["source_module"])
        if adapter:
            qs = qs.filter(source_content_type=adapter.content_type())
    if data.get("risk_level"):
        qs = qs.filter(hazards__initial_risk_level=data["risk_level"])
    if data.get("from_date"):
        qs = qs.filter(assessment_date__gte=data["from_date"])
    if data.get("to_date"):
        qs = qs.filter(assessment_date__lte=data["to_date"])
    financial_year = data.get("financial_year")
    if financial_year:
        try:
            start_year = int(str(financial_year)[:4])
        except (TypeError, ValueError):
            start_year = None
        if start_year:
            qs = qs.filter(
                assessment_date__gte=date(start_year, 4, 1),
                assessment_date__lte=date(start_year + 1, 3, 31),
            )
    return qs.distinct()


def _form_data(request, form=None):
    return form.cleaned_data if form and form.is_valid() else request.GET


def filter_source_queryset(request, form=None):
    data = _form_data(request, form)
    adapter = get_adapter(data.get("source_module"))
    if not adapter:
        return None, None
    return adapter, adapter.filtered_queryset(data, user=request.user)


def _linked_hazards_for_source(adapter, records, data=None):
    content_type = adapter.content_type()
    object_ids = [record.pk for record in records]
    hazards = HIRAHazard.objects.select_related(
        "hira",
        "hira__plant",
        "hira__department",
        "hira__module",
        "responsible_person",
    ).prefetch_related("actions").filter(
        Q(related_content_type=content_type, related_object_id__in=object_ids)
        | Q(hira__source_content_type=content_type, hira__source_object_id__in=object_ids)
    ).distinct()
    if data and data.get("risk_level"):
        hazards = hazards.filter(initial_risk_level=data["risk_level"])
    if data and data.get("module"):
        hazards = hazards.filter(hira__module=data["module"])
    return list(hazards)


def build_source_first_hira_workbook(adapter, records, data=None):
    records = list(records)
    linked_hazards = _linked_hazards_for_source(adapter, records, data=data)
    hazards_by_source = {}
    content_type = adapter.content_type()
    for hazard in linked_hazards:
        key = (
            hazard.related_object_id
            if hazard.related_content_type_id == content_type.id and hazard.related_object_id
            else hazard.hira.source_object_id
        )
        hazards_by_source.setdefault(key, []).append(hazard)

    hiras = list({hazard.hira for hazard in linked_hazards})
    actions = list(
        HIRAAction.objects.select_related(
            "hazard",
            "hazard__hira",
            "responsible_person",
            "verified_by",
        ).filter(hazard__in=linked_hazards)
    )

    wb = Workbook()
    wb.remove(wb.active)

    level_counts = Counter(h.initial_risk_level for h in linked_hazards)
    residual_counts = Counter(h.residual_risk_level for h in linked_hazards)
    action_counts = Counter(a.status for a in actions)

    summary = wb.create_sheet("HIRA Dashboard")
    summary.append(["EHS-360 | REAL-TIME HIRA DASHBOARD"])
    summary.append([])
    summary.append(["Metric", "Value", None, "Risk Level", "Initial", "Residual"])
    summary_rows = [
        ("Source Module", adapter.module_name),
        ("Current Source Records", len(records)),
        ("Linked HIRA Assessments", len(hiras)),
        ("Assessed Hazard Rows", len(linked_hazards)),
        ("Pending Assessment Records", len([r for r in records if r.pk not in hazards_by_source])),
        ("Critical Initial Risks", level_counts["Critical"]),
        ("High Initial Risks", level_counts["High"]),
        ("Medium Initial Risks", level_counts["Medium"]),
        ("Low Initial Risks", level_counts["Low"]),
        ("Open / In Progress Actions", action_counts["OPEN"] + action_counts["IN_PROGRESS"]),
        ("Completed / Closed Actions", action_counts["COMPLETED"] + action_counts["VERIFIED"] + action_counts["CLOSED"]),
    ]
    for idx, row in enumerate(summary_rows, start=4):
        summary.cell(row=idx, column=1, value=row[0])
        summary.cell(row=idx, column=2, value=row[1])
    for offset, level in enumerate(["Critical", "High", "Medium", "Low"], start=4):
        summary.cell(row=offset, column=4, value=level)
        summary.cell(row=offset, column=5, value=level_counts[level])
        summary.cell(row=offset, column=6, value=residual_counts[level])
    if linked_hazards:
        chart = BarChart()
        chart.title = "Risk Distribution"
        chart.y_axis.title = "Count"
        data_ref = Reference(summary, min_col=5, max_col=6, min_row=3, max_row=7)
        cats = Reference(summary, min_col=4, min_row=4, max_row=7)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        summary.add_chart(chart, "H4")
    _style_sheet(summary)
    summary["A1"].font = Font(bold=True, size=16, color="1F4E78")
    summary["A1"].fill = PatternFill("solid", fgColor="FFFFFF")
    _autosize(summary)

    register = wb.create_sheet("HIRA Register")
    headers = [
        "HIRA ID",
        "Plant",
        "Department",
        "Module",
        "Source Module",
        "Source Record",
        "Source Date",
        "Source Status",
        "Process",
        "Activity",
        "Routine / Non-Routine",
        "Hazard Category",
        "Hazard",
        "Unsafe Act",
        "Unsafe Condition",
        "Potential Consequence",
        "Affected Persons",
        "Existing Source Controls / Actions",
        "Existing HIRA Controls",
        "Control Type",
        "Likelihood (L)",
        "Severity (S)",
        "Initial Risk",
        "Initial Risk Level",
        "Additional Controls / Recommendations",
        "Action Required",
        "Responsible Person",
        "Target Date",
        "Priority",
        "Action Status",
        "Residual Likelihood",
        "Residual Severity",
        "Residual Risk",
        "Residual Risk Level",
        "Review Date",
        "HIRA Status",
        "Evidence / Remarks",
    ]
    register.append(headers)

    for record in records:
        source_info = adapter.source_info(record)
        linked = hazards_by_source.get(record.pk, [])
        if not linked:
            register.append([
                "",
                source_info.plant.name if source_info.plant else "",
                source_info.department.name if source_info.department else "",
                adapter.module_name,
                source_info.module_name,
                source_info.record_label,
                source_info.record_date,
                source_info.status,
                source_info.process,
                source_info.activity,
                "",
                "",
                "",
                source_info.unsafe_act,
                source_info.unsafe_condition,
                "",
                "",
                source_info.existing_action,
                "",
                "",
                "",
                "",
                "",
                "Not Assessed",
                "",
                "No",
                source_info.responsible_person.get_full_name() if source_info.responsible_person else "",
                source_info.target_date,
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "Not Assessed",
                "",
            ])
            continue

        for hazard in linked:
            first_action = hazard.actions.first()
            register.append([
                hazard.hira.hira_number,
                source_info.plant.name if source_info.plant else hazard.hira.plant.name,
                source_info.department.name if source_info.department else hazard.hira.department.name,
                hazard.hira.module.name,
                source_info.module_name,
                source_info.record_label,
                source_info.record_date,
                source_info.status,
                source_info.process or hazard.hira.process,
                hazard.activity or source_info.activity,
                hazard.hira.get_assessment_type_display(),
                hazard.hazard_category,
                hazard.hazard,
                hazard.unsafe_act or source_info.unsafe_act,
                hazard.unsafe_condition or source_info.unsafe_condition,
                hazard.potential_consequence,
                hazard.persons_exposed,
                source_info.existing_action,
                hazard.existing_controls,
                hazard.control_hierarchy,
                hazard.likelihood,
                hazard.severity,
                hazard.initial_risk_score,
                hazard.initial_risk_level,
                hazard.additional_controls,
                "Yes" if hazard.action_required else "No",
                hazard.responsible_person.get_full_name() if hazard.responsible_person else "",
                hazard.target_date,
                first_action.priority if first_action else hazard.initial_risk_level,
                _action_status_text(first_action.status) if first_action else "",
                hazard.residual_likelihood,
                hazard.residual_severity,
                hazard.residual_risk_score,
                hazard.residual_risk_level,
                hazard.hira.review_date,
                _status_text(hazard.hira.status),
                hazard.evidence_remarks,
            ])

    if not records:
        register.append([f"No {adapter.module_name} records found for the selected filters."])
    _style_sheet(register)
    for row in register.iter_rows(min_row=2):
        for cell in row:
            fill = _risk_fill(cell.value)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
    _autosize(register, 35)

    _append_reference_sheets(wb, actions, hiras)
    return wb


def _append_reference_sheets(wb, actions=None, hiras=None):
    actions = actions or []
    hiras = hiras or []

    matrix = wb.create_sheet("Risk Matrix")
    matrix.append(["Likelihood / Severity", "1 - Insignificant", "2 - Minor", "3 - Moderate", "4 - Major", "5 - Catastrophic"])
    for likelihood, label in enumerate(["Rare", "Unlikely", "Possible", "Likely", "Almost Certain"], start=1):
        matrix.append([f"{likelihood} - {label}", *[likelihood * severity for severity in range(1, 6)]])
    _style_sheet(matrix)
    _autosize(matrix)

    levels = wb.create_sheet("Risk Levels")
    levels.append(["Risk Score", "Risk Level", "Meaning", "Recommended Action"])
    for row in _configured_risk_level_rows():
        levels.append(row)
    _style_sheet(levels)
    _autosize(levels)

    controls = wb.create_sheet("Hierarchy of Controls")
    controls.append(["Priority", "Control Type", "Description", "Example"])
    for row in [
        (1, "Elimination", "Remove the hazard completely", "Eliminate work at height by doing work at ground level"),
        (2, "Substitution", "Replace with a less hazardous option", "Use a less hazardous chemical"),
        (3, "Engineering Control", "Physically isolate people from the hazard", "Machine guard, interlock, ventilation"),
        (4, "Administrative Control", "Change the way work is performed", "SOP, training, permit, signage"),
        (5, "PPE", "Protect the worker with personal protective equipment", "Helmet, gloves, goggles, respirator"),
    ]:
        controls.append(row)
    _style_sheet(controls)
    _autosize(controls)

    action_ws = wb.create_sheet("Action Register")
    action_ws.append(["Action ID", "HIRA ID", "Source Record", "Activity", "Risk / Hazard", "Action / Recommendation", "Responsible Person", "Target Date", "Priority", "Status", "Completion Date", "Verification", "Evidence / Closure Remarks"])
    for action in actions:
        source_info = describe_source_object(action.hazard.related_record or action.hazard.hira.source_record)
        action_ws.append([
            action.action_id,
            action.hazard.hira.hira_number,
            source_info.record_label if source_info else "",
            action.hazard.activity,
            action.hazard.hazard,
            action.action_description,
            action.responsible_person.get_full_name() if action.responsible_person else "",
            action.target_date,
            action.priority,
            _action_status_text(action.status),
            action.completion_date,
            action.verification,
            action.remarks,
        ])
    _style_sheet(action_ws)
    _autosize(action_ws)

    approval = wb.create_sheet("Approval & Revision")
    approval.append(["HIRA Document Information", "Value", "Prepared By", "Reviewed By", "Approved By"])
    for hira in hiras:
        approval.append([
            hira.hira_number,
            f"{hira.plant.name} | {hira.department.name} | Rev {hira.revision_number}",
            hira.prepared_by.get_full_name() if hira.prepared_by else "",
            hira.reviewed_by.get_full_name() if hira.reviewed_by else "",
            hira.approved_by.get_full_name() if hira.approved_by else "",
        ])
    _style_sheet(approval)
    _autosize(approval)

    review = wb.create_sheet("Review Triggers")
    review.append(["Trigger", "When HIRA Must Be Reviewed / Reassessed", "Example"])
    for row in [
        ("New process / activity", "Before introducing a new activity", "New production line"),
        ("New machinery", "Before commissioning or operation", "New CNC machine / press"),
        ("Chemical change", "Before use of new or changed chemical", "New solvent"),
        ("Process change", "When process, layout or operating conditions change", "Production capacity increase"),
        ("Incident / near miss", "After an event indicates controls may be inadequate", "Machine injury / near miss"),
        ("Control failure", "When an existing control is ineffective", "Guard/interlock failure"),
        ("Legal / regulatory change", "When applicable requirements change", "New statutory requirement"),
    ]:
        review.append(row)
    _style_sheet(review)
    _autosize(review)

    masters = wb.create_sheet("Industry Activity Master")
    masters.append(["Industry / Area", "Common Process", "Typical Activities"])
    for row in [
        ("Training Management", "Safety Training", "Training planning; preparation; delivery; practical demonstration; assessment; closure"),
        ("Manufacturing", "Production", "Machine operation; cutting; pressing; assembly; grinding"),
        ("Engineering", "Fabrication", "Welding; gas cutting; grinding; drilling; machining"),
        ("Warehouse & Logistics", "Material Movement", "Forklift; loading/unloading; stacking; manual handling"),
        ("Maintenance", "Plant Maintenance", "Electrical work; LOTO; work at height; confined space"),
    ]:
        masters.append(row)
    _style_sheet(masters)
    _autosize(masters)

    hazard_master = wb.create_sheet("Hazard Master")
    hazard_master.append(["Hazard Category", "Typical Hazards", "Typical Consequences"])
    for row in [
        ("Mechanical", "Moving parts; pinch points; sharp edges; stored energy", "Cuts; crush injury; amputation"),
        ("Electrical", "Live parts; arc flash; damaged cables; poor isolation", "Shock; burns; fatality; fire"),
        ("Chemical", "Toxic; corrosive; flammable; reactive; splash", "Burns; poisoning; inhalation; fire"),
        ("Physical", "Noise; heat; cold; radiation; vibration", "Hearing loss; burns; heat stress; exposure"),
        ("Fire / Explosion", "Ignition sources; combustible materials; gas release", "Fire; explosion; burns; fatality"),
        ("Ergonomic", "Manual handling; repetitive motion; poor posture", "MSD; strains; fatigue"),
        ("Vehicle", "Mobile equipment; reversing; pedestrian interaction", "Collision; crush injury; fatality"),
    ]:
        hazard_master.append(row)
    _style_sheet(hazard_master)
    _autosize(hazard_master)


def build_hira_workbook(hiras, filters=None):
    hiras = list(hiras)
    hazards = list(
        HIRAHazard.objects.select_related(
            "hira",
            "hira__plant",
            "hira__department",
            "hira__module",
            "hira__source_content_type",
            "responsible_person",
            "related_content_type",
        ).filter(hira__in=hiras)
    )
    actions = list(
        HIRAAction.objects.select_related(
            "hazard",
            "hazard__hira",
            "responsible_person",
            "verified_by",
        ).filter(hazard__hira__in=hiras)
    )

    wb = Workbook()
    wb.remove(wb.active)

    summary = wb.create_sheet("HIRA Dashboard")
    summary.append(["EHS-360 | INDUSTRIAL HIRA DASHBOARD"])
    summary.append([])
    level_counts = Counter(h.initial_risk_level for h in hazards)
    residual_counts = Counter(h.residual_risk_level for h in hazards)
    action_counts = Counter(a.status for a in actions)
    summary_rows = [
        ("Total HIRA Assessments", len(hiras)),
        ("Total Activities", len({h.activity for h in hazards})),
        ("Total Hazards", len(hazards)),
        ("Critical Initial Risks", level_counts["Critical"]),
        ("High Initial Risks", level_counts["High"]),
        ("Medium Initial Risks", level_counts["Medium"]),
        ("Low Initial Risks", level_counts["Low"]),
        ("Open / In Progress Actions", action_counts["OPEN"] + action_counts["IN_PROGRESS"]),
        ("Completed / Closed Actions", action_counts["COMPLETED"] + action_counts["VERIFIED"] + action_counts["CLOSED"]),
    ]
    summary.append(["Metric", "Value", None, "Risk Level", "Initial", "Residual"])
    for idx, row in enumerate(summary_rows, start=4):
        summary.cell(row=idx, column=1, value=row[0])
        summary.cell(row=idx, column=2, value=row[1])
    for offset, level in enumerate(["Critical", "High", "Medium", "Low"], start=4):
        summary.cell(row=offset, column=4, value=level)
        summary.cell(row=offset, column=5, value=level_counts[level])
        summary.cell(row=offset, column=6, value=residual_counts[level])
    if hazards:
        chart = BarChart()
        chart.title = "Risk Distribution"
        chart.y_axis.title = "Count"
        data = Reference(summary, min_col=5, max_col=6, min_row=3, max_row=7)
        cats = Reference(summary, min_col=4, min_row=4, max_row=7)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        summary.add_chart(chart, "H4")
    _style_sheet(summary)
    summary["A1"].font = Font(bold=True, size=16, color="1F4E78")
    summary["A1"].fill = PatternFill("solid", fgColor="FFFFFF")
    _autosize(summary)

    register = wb.create_sheet("HIRA Register")
    headers = [
        "HIRA ID",
        "Plant",
        "Department",
        "Module",
        "Source Module",
        "Source Record",
        "Source Date",
        "Source Status",
        "Process",
        "Activity",
        "Routine / Non-Routine",
        "Hazard Category",
        "Hazard",
        "Unsafe Act",
        "Unsafe Condition",
        "Potential Consequence",
        "Affected Persons",
        "Existing Controls",
        "Control Type",
        "Likelihood (L)",
        "Severity (S)",
        "Initial Risk",
        "Initial Risk Level",
        "Additional Controls / Recommendations",
        "Action Required",
        "Responsible Person",
        "Target Date",
        "Priority",
        "Action Status",
        "Residual Likelihood",
        "Residual Severity",
        "Residual Risk",
        "Residual Risk Level",
        "Related Record",
        "Review Date",
        "Status",
        "Evidence / Remarks",
    ]
    register.append(headers)
    for hazard in hazards:
        first_action = hazard.actions.first()
        source_info = describe_source_object(hazard.related_record or hazard.hira.source_record)
        register.append([
            hazard.hira.hira_number,
            hazard.hira.plant.name,
            hazard.hira.department.name,
            hazard.hira.module.name,
            source_info.module_name if source_info else "",
            source_info.record_label if source_info else "",
            source_info.record_date if source_info else "",
            source_info.status if source_info else "",
            hazard.hira.process,
            hazard.activity,
            hazard.hira.get_assessment_type_display(),
            hazard.hazard_category,
            hazard.hazard,
            hazard.unsafe_act,
            hazard.unsafe_condition,
            hazard.potential_consequence,
            hazard.persons_exposed,
            hazard.existing_controls,
            hazard.control_hierarchy,
            hazard.likelihood,
            hazard.severity,
            hazard.initial_risk_score,
            hazard.initial_risk_level,
            hazard.additional_controls,
            "Yes" if hazard.action_required else "No",
            hazard.responsible_person.get_full_name() if hazard.responsible_person else "",
            hazard.target_date,
            first_action.priority if first_action else hazard.initial_risk_level,
            _action_status_text(first_action.status) if first_action else "",
            hazard.residual_likelihood,
            hazard.residual_severity,
            hazard.residual_risk_score,
            hazard.residual_risk_level,
            source_info.record_label if source_info else "",
            hazard.hira.review_date,
            _status_text(hazard.hira.status),
            hazard.evidence_remarks,
        ])
    if not hazards:
        register.append(["No HIRA records found for the selected filters."])
    _style_sheet(register)
    for row in register.iter_rows(min_row=2):
        for cell in row:
            fill = _risk_fill(cell.value)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
    _autosize(register, 35)

    matrix = wb.create_sheet("Risk Matrix")
    matrix.append(["Likelihood / Severity", "1 - Insignificant", "2 - Minor", "3 - Moderate", "4 - Major", "5 - Catastrophic"])
    for likelihood, label in enumerate(["Rare", "Unlikely", "Possible", "Likely", "Almost Certain"], start=1):
        matrix.append([f"{likelihood} - {label}", *[likelihood * severity for severity in range(1, 6)]])
    _style_sheet(matrix)
    _autosize(matrix)

    levels = wb.create_sheet("Risk Levels")
    levels.append(["Risk Score", "Risk Level", "Meaning", "Recommended Action"])
    for row in _configured_risk_level_rows():
        levels.append(row)
    _style_sheet(levels)
    _autosize(levels)

    controls = wb.create_sheet("Hierarchy of Controls")
    controls.append(["Priority", "Control Type", "Description", "Example"])
    controls_rows = [
        (1, "Elimination", "Remove the hazard completely", "Eliminate work at height by doing work at ground level"),
        (2, "Substitution", "Replace with a less hazardous option", "Use a less hazardous chemical"),
        (3, "Engineering Control", "Physically isolate people from the hazard", "Machine guard, interlock, ventilation"),
        (4, "Administrative Control", "Change the way work is performed", "SOP, training, permit, signage"),
        (5, "PPE", "Protect the worker with personal protective equipment", "Helmet, gloves, goggles, respirator"),
    ]
    for row in controls_rows:
        controls.append(row)
    _style_sheet(controls)
    _autosize(controls)

    action_ws = wb.create_sheet("Action Register")
    action_ws.append(["Action ID", "HIRA ID", "Source Record", "Activity", "Risk / Hazard", "Action / Recommendation", "Responsible Person", "Target Date", "Priority", "Status", "Completion Date", "Verification", "Evidence / Closure Remarks"])
    for action in actions:
        source_info = describe_source_object(action.hazard.related_record or action.hazard.hira.source_record)
        action_ws.append([
            action.action_id,
            action.hazard.hira.hira_number,
            source_info.record_label if source_info else "",
            action.hazard.activity,
            action.hazard.hazard,
            action.action_description,
            action.responsible_person.get_full_name() if action.responsible_person else "",
            action.target_date,
            action.priority,
            _action_status_text(action.status),
            action.completion_date,
            action.verification,
            action.remarks,
        ])
    _style_sheet(action_ws)
    _autosize(action_ws)

    approval = wb.create_sheet("Approval & Revision")
    approval.append(["HIRA Document Information", "Value", "Prepared By", "Reviewed By", "Approved By"])
    for hira in hiras:
        approval.append([
            hira.hira_number,
            f"{hira.plant.name} | {hira.department.name} | Rev {hira.revision_number}",
            hira.prepared_by.get_full_name() if hira.prepared_by else "",
            hira.reviewed_by.get_full_name() if hira.reviewed_by else "",
            hira.approved_by.get_full_name() if hira.approved_by else "",
        ])
    _style_sheet(approval)
    _autosize(approval)

    review = wb.create_sheet("Review Triggers")
    review.append(["Trigger", "When HIRA Must Be Reviewed / Reassessed", "Example"])
    for row in [
        ("New process / activity", "Before introducing a new activity", "New production line"),
        ("New machinery", "Before commissioning or operation", "New CNC machine / press"),
        ("Chemical change", "Before use of new or changed chemical", "New solvent"),
        ("Process change", "When process, layout or operating conditions change", "Production capacity increase"),
        ("Incident / near miss", "After an event indicates controls may be inadequate", "Machine injury / near miss"),
        ("Control failure", "When an existing control is ineffective", "Guard/interlock failure"),
        ("Legal / regulatory change", "When applicable requirements change", "New statutory requirement"),
    ]:
        review.append(row)
    _style_sheet(review)
    _autosize(review)

    masters = wb.create_sheet("Industry Activity Master")
    masters.append(["Industry / Area", "Common Process", "Typical Activities"])
    for row in [
        ("Training Management", "Safety Training", "Training planning; preparation; delivery; practical demonstration; assessment; closure"),
        ("Manufacturing", "Production", "Machine operation; cutting; pressing; assembly; grinding"),
        ("Engineering", "Fabrication", "Welding; gas cutting; grinding; drilling; machining"),
        ("Warehouse & Logistics", "Material Movement", "Forklift; loading/unloading; stacking; manual handling"),
        ("Maintenance", "Plant Maintenance", "Electrical work; LOTO; work at height; confined space"),
    ]:
        masters.append(row)
    _style_sheet(masters)
    _autosize(masters)

    hazard_master = wb.create_sheet("Hazard Master")
    hazard_master.append(["Hazard Category", "Typical Hazards", "Typical Consequences"])
    for row in [
        ("Mechanical", "Moving parts; pinch points; sharp edges; stored energy", "Cuts; crush injury; amputation"),
        ("Electrical", "Live parts; arc flash; damaged cables; poor isolation", "Shock; burns; fatality; fire"),
        ("Chemical", "Toxic; corrosive; flammable; reactive; splash", "Burns; poisoning; inhalation; fire"),
        ("Physical", "Noise; heat; cold; radiation; vibration", "Hearing loss; burns; heat stress; exposure"),
        ("Fire / Explosion", "Ignition sources; combustible materials; gas release", "Fire; explosion; burns; fatality"),
        ("Ergonomic", "Manual handling; repetitive motion; poor posture", "MSD; strains; fatigue"),
        ("Vehicle", "Mobile equipment; reversing; pedestrian interaction", "Collision; crush injury; fatality"),
    ]:
        hazard_master.append(row)
    _style_sheet(hazard_master)
    _autosize(hazard_master)

    return wb


def _xlsx_response(wb, filename_prefix="HIRA_Report"):
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    response["Content-Disposition"] = f'attachment; filename="{filename_prefix}_{stamp}.xlsx"'
    return response


def workbook_response(hiras):
    return _xlsx_response(build_hira_workbook(hiras))


def source_first_workbook_response(adapter, records, data=None):
    slug = adapter.module_name.replace(" ", "_")
    return _xlsx_response(
        build_source_first_hira_workbook(adapter, records, data=data),
        filename_prefix=f"HIRA_{slug}_Realtime_Report",
    )
