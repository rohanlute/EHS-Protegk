# apps/reports/services/excel_engine.py
"""
Shared Excel engine used by every EHS module.

A ReportSpec declares WHAT to render; this engine decides HOW.
"""
from dataclasses import dataclass, field
from io import BytesIO
from datetime import datetime
from typing import Callable, Optional

from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


# ──────────────────────────────────────────────────────────────
# STYLE CONSTANTS
# ──────────────────────────────────────────────────────────────
class Style:
    HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    TITLE_FONT  = Font(bold=True, size=16, color="1F4E78")
    SECTION_FONT = Font(bold=True, size=12, color="1F4E78")
    BOLD = Font(bold=True)
    ITALIC = Font(italic=True, color="666666")
    THIN = Side(style="thin", color="D0D0D0")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    CENTER = Alignment(horizontal="center", vertical="center")
    WRAP = Alignment(vertical="top", wrap_text=True)

    LEVEL_FILLS = {
        "Critical": PatternFill("solid", fgColor="FFC7CE"),
        "High":     PatternFill("solid", fgColor="FFD9B3"),
        "Medium":   PatternFill("solid", fgColor="FFEB9C"),
        "Low":      PatternFill("solid", fgColor="C6EFCE"),
        "Pass":     PatternFill("solid", fgColor="C6EFCE"),
        "Fail":     PatternFill("solid", fgColor="FFC7CE"),
        "NA":       PatternFill("solid", fgColor="E7E6E6"),
        "Open":     PatternFill("solid", fgColor="FFEB9C"),
        "In Progress": PatternFill("solid", fgColor="FFD9B3"),
        "Completed": PatternFill("solid", fgColor="C6EFCE"),
        "Closed":   PatternFill("solid", fgColor="C6EFCE"),
        "Overdue":  PatternFill("solid", fgColor="FFC7CE"),
        "Converted": PatternFill("solid", fgColor="C6EFCE"),
        "Assigned": PatternFill("solid", fgColor="FFEB9C"),
    }


# ──────────────────────────────────────────────────────────────
# SPEC DATACLASSES
# ──────────────────────────────────────────────────────────────
@dataclass
class Column:
    """One column in a register/table sheet."""
    header: str
    value: Optional[Callable] = None
    formula: Optional[Callable] = None
    width: int = 18
    fill_by_value: bool = False
    bold: bool = False
    wrap: bool = True


@dataclass
class KPI:
    label: str
    formula: str
    fill_color: Optional[str] = None


@dataclass
class Pivot:
    title: str
    row_header: str
    row_values: list
    col_header: str
    col_formula_template: str


@dataclass
class StaticTable:
    name: str
    headers: list
    rows: list
    widths: Optional[list] = None


@dataclass
class RegisterSheet:
    name: str
    columns: list
    queryset: object
    extend_formulas_to: int = 500
    dropdowns: dict = field(default_factory=dict)


@dataclass
class DashboardSheet:
    name: str = "Dashboard"
    title: str = ""
    kpis: list = field(default_factory=list)
    pivots: list = field(default_factory=list)
    footer_note: str = ""


@dataclass
class ReportSpec:
    document_title: str
    file_prefix: str
    sheets: list


# ──────────────────────────────────────────────────────────────
# RENDERERS
# ──────────────────────────────────────────────────────────────
def _style_header_row(ws, row=1):
    for cell in ws[row]:
        cell.fill = Style.HEADER_FILL
        cell.font = Style.HEADER_FONT
        cell.alignment = Style.CENTER
        cell.border = Style.BORDER
    ws.freeze_panes = f"A{row + 1}"


def _autosize(ws, max_w=45):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letter].width = min(longest + 2, max_w)


def _render_dashboard(wb, spec: DashboardSheet):
    ws = wb.create_sheet(spec.name)
    if spec.title:
        ws["A1"] = spec.title
        ws["A1"].font = Style.TITLE_FONT

    ws["A3"], ws["B3"] = "Metric", "Value"
    ws["A3"].font = Style.BOLD
    ws["B3"].font = Style.BOLD
    for i, kpi in enumerate(spec.kpis, start=4):
        ws.cell(row=i, column=1, value=kpi.label).font = Style.BOLD
        ws.cell(row=i, column=1).fill = (
            PatternFill("solid", fgColor=kpi.fill_color)
            if kpi.fill_color else PatternFill("solid", fgColor="EAF1F8")
        )
        ws.cell(row=i, column=1).border = Style.BORDER
        ws.cell(row=i, column=2, value=kpi.formula).border = Style.BORDER

    col = 4
    for pivot in spec.pivots:
        ws.cell(row=3, column=col, value=pivot.title).font = Style.SECTION_FONT
        ws.cell(row=3, column=col + 1, value=pivot.row_header).font = Style.BOLD
        ws.cell(row=3, column=col + 2, value=pivot.col_header).font = Style.BOLD

        for j, val in enumerate(pivot.row_values, start=4):
            ws.cell(row=j, column=col, value=val).font = Style.BOLD
            fill = Style.LEVEL_FILLS.get(val)
            if fill:
                ws.cell(row=j, column=col).fill = fill
            formula = pivot.col_formula_template.format(value=val)
            ws.cell(row=j, column=col + 1, value=formula)
            ws.cell(row=j, column=col + 2, value=formula)
        col += 4

    if spec.footer_note:
        last = max(ws.max_row + 2, 20)
        ws.cell(row=last, column=1, value=spec.footer_note).font = Style.ITALIC

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 20


def _render_register(wb, spec: RegisterSheet):
    ws = wb.create_sheet(spec.name)
    ws.append([c.header for c in spec.columns])
    _style_header_row(ws)

    row = 2
    for obj in spec.queryset:
        for idx, col in enumerate(spec.columns, start=1):
            if col.formula:
                val = col.formula(obj, row)
            else:
                val = col.value(obj) if col.value else ""
            cell = ws.cell(row=row, column=idx, value=val)
            if col.wrap:
                cell.alignment = Style.WRAP
            if col.fill_by_value:
                fill = Style.LEVEL_FILLS.get(str(cell.value))
                if fill:
                    cell.fill = fill
                    cell.font = Style.BOLD
            if col.bold:
                cell.font = Style.BOLD
        row += 1

    last_data_row = row - 1
    if spec.extend_formulas_to and last_data_row < spec.extend_formulas_to:
        for r in range(last_data_row + 1, spec.extend_formulas_to + 1):
            for idx, col in enumerate(spec.columns, start=1):
                if col.formula:
                    ws.cell(row=r, column=idx, value=col.formula(None, r))

    for idx, col in enumerate(spec.columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = col.width

    for col_idx, options in spec.dropdowns.items():
        letter = get_column_letter(col_idx)
        dv = DataValidation(
            type="list",
            formula1=f'"{options}"',
            allow_blank=True,
        )
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{spec.extend_formulas_to}")


def _render_static(wb, spec: StaticTable):
    ws = wb.create_sheet(spec.name)
    ws.append(spec.headers)
    _style_header_row(ws)
    for r in spec.rows:
        ws.append(list(r))
    _autosize(ws)
    if spec.widths:
        for idx, w in enumerate(spec.widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w


# ──────────────────────────────────────────────────────────────
# MAIN ENTRY POINTS
# ──────────────────────────────────────────────────────────────
def build_workbook(spec: ReportSpec) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    for sheet in spec.sheets:
        if isinstance(sheet, DashboardSheet):
            _render_dashboard(wb, sheet)
        elif isinstance(sheet, RegisterSheet):
            _render_register(wb, sheet)
        elif isinstance(sheet, StaticTable):
            _render_static(wb, sheet)
        else:
            raise TypeError(f"Unknown sheet type: {type(sheet)}")

    return wb


def workbook_http_response(spec: ReportSpec) -> HttpResponse:
    wb = build_workbook(spec)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{spec.file_prefix}_{stamp}.xlsx"'
    )
    return response